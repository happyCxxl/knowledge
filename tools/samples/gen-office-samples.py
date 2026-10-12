#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成 Office 解析器的复核样例（对应 docs/stages/解析环节/05-待办清单.md 第四节 OFF-01~OFF-14）。

样例落在 samples/office-parse-NN/（该目录不入库），每条目一个目录、文件名带序号与要点。
OFF-12 / OFF-13 需要写入公式缓存值，openpyxl 只写公式不写缓存，改由 tools/samples/gen-office-poi/ 下的
Java 小程序用 POI 生成；OFF-15（.doc 域代码）POI 无法写出旧二进制格式，目录内只放说明。

用法：
    <python> tools/samples/gen-office-samples.py
"""

import base64
import os
import struct
import zlib

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pptx import Presentation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SAMPLES = os.path.join(ROOT, "samples")


def out_dir(item: str) -> str:
    """条目目录：samples/office-parse-NN"""
    path = os.path.join(SAMPLES, "office-parse-%s" % item)
    os.makedirs(path, exist_ok=True)
    return path


def tiny_png() -> bytes:
    """1×1 透明 PNG（手写字节，避免依赖 Pillow 版本差异）"""
    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    header = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    raw = b"\x00\x00\x00\x00\x00"
    return header + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


# ---------------------------------------------------------------- 逐条目生成

def gen_01_pptx():
    """OFF-01：一页含文本框、2×2 表格与一张图片（单元 = 幻灯片）"""
    directory = out_dir("01")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Pt(60), Pt(40), Pt(400), Pt(40))
    box.text_frame.text = "季度汇报"
    table = slide.shapes.add_table(2, 2, Pt(60), Pt(120), Pt(400), Pt(80)).table
    for r in range(2):
        for c in range(2):
            table.cell(r, c).text = ["评分项", "分值", "A1 报价", "30"][r * 2 + c]
    slide.shapes.add_picture(_png_file(directory), Pt(60), Pt(240), Pt(60), Pt(60))
    prs.save(os.path.join(directory, "01-一页含文本框表格与图片.pptx"))


def _png_file(directory: str) -> str:
    """把样例 PNG 落盘供 python-pptx 引用（文件名固定，重复生成覆盖）"""
    path = os.path.join(directory, "sample.png")
    with open(path, "wb") as handle:
        handle.write(tiny_png())
    return path


def gen_02_docx():
    """OFF-02：三段，中间一段为空（段落序号应取源文档序号 0/2）"""
    directory = out_dir("02")
    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph()
    doc.add_paragraph("第三段")
    doc.save(os.path.join(directory, "01-中间空段.docx"))


def gen_03_docx():
    """OFF-03：只有图片没有文字 / 图片加长正文（疑似图片单元判定两条口径）"""
    directory = out_dir("03")
    image = _png_file(directory)

    only_image = Document()
    only_image.add_picture(image, width=Pt(60))
    only_image.save(os.path.join(directory, "01-只有一张图无文字.docx"))

    with_text = Document()
    for i in range(30):
        with_text.add_paragraph("这是第 %d 段正文，用于把字符数抬到门槛之上。" % (i + 1))
    with_text.add_picture(image, width=Pt(60))
    with_text.save(os.path.join(directory, "02-有图但正文很多.docx"))


def gen_04_docx():
    """OFF-04：长正文（跨多个 3000 字符预算单元）后跟一张图片（图片应归最后一个单元）"""
    directory = out_dir("04")
    doc = Document()
    for i in range(40):
        doc.add_paragraph(("第 %d 段：" % (i + 1)) + "正文内容。" * 60)
    doc.add_picture(_png_file(directory), width=Pt(60))
    doc.save(os.path.join(directory, "01-长正文后带图片.docx"))


def gen_05_docx():
    """OFF-05：只有表格（单元格文本应计入预算）/ 只有页眉（页眉文本应计入预算）"""
    directory = out_dir("05")

    table_only = Document()
    table = table_only.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "abcdef"
    table.cell(0, 1).text = "ghijkl"
    table_only.save(os.path.join(directory, "01-只有表格.docx"))

    header_only = Document()
    header = header_only.sections[0].header
    header.paragraphs[0].text = "页眉文字"
    header_only.save(os.path.join(directory, "02-只有页眉.docx"))


def gen_06_xlsx():
    """OFF-06：表头行指向真实数据首行（首行空 → 1；表头在首行 → 0）"""
    directory = out_dir("06")

    wb = Workbook()
    ws = wb.active
    ws.title = "首行空"
    ws["A1"] = None
    for col, value in enumerate(["评分项", "分值", "备注"], start=1):
        ws.cell(row=2, column=col, value=value)
    for col, value in enumerate(["A1 报价", 30, "首行为空"], start=1):
        ws.cell(row=3, column=col, value=value)
    wb.save(os.path.join(directory, "01-首行空数据从第二行.xlsx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "表头在首行"
    for col, value in enumerate(["评分项", "分值"], start=1):
        ws.cell(row=1, column=col, value=value)
    for col, value in enumerate(["A1 报价", 30], start=1):
        ws.cell(row=2, column=col, value=value)
    wb.save(os.path.join(directory, "02-表头在首行.xlsx"))


def gen_07_xlsx():
    """OFF-07：首行空三列（cols=3）/ 窄标题行四列（cols 不被首行截短）"""
    directory = out_dir("07")

    wb = Workbook()
    ws = wb.active
    ws.title = "空首行三列"
    ws.cell(row=1, column=1, value=None)
    for col, value in enumerate(["评分项", "分值", "说明"], start=1):
        ws.cell(row=2, column=col, value=value)
    for col, value in enumerate(["A1 报价", 30, "三列数据"], start=1):
        ws.cell(row=3, column=col, value=value)
    wb.save(os.path.join(directory, "01-首行空三列.xlsx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "窄标题"
    ws.cell(row=1, column=1, value="导出报表")
    for col, value in enumerate(["A", "B", "C", "D"], start=1):
        ws.cell(row=2, column=col, value=value)
    wb.save(os.path.join(directory, "02-窄标题行四列.xlsx"))


def gen_08_samples():
    """OFF-08：合并覆盖格 / vMerge 续接格的文本样例由 POI 生成（见 tools/samples/gen-office-poi/）

    openpyxl 不允许给被合并覆盖格写值（MergedCell 只读）、python-docx 写 w:vMerge 的元素顺序不可控，
    这两类样例统一交给 POI 小程序，保证与解析器读到的结构一致。
    """
    directory = out_dir("08")
    with open(os.path.join(directory, "说明-本目录样例由 POI 生成.md"), "w", encoding="utf-8") as handle:
        handle.write(
            "# OFF-08 样例说明\n\n"
            "本目录两份样例（横向合并被覆盖格带文本.xlsx、纵向合并续接格带文本.docx）由\n"
            "`tools/samples/gen-office-poi/GenOfficePoiSamples.java` 用 POI 生成：\n\n"
            "- 合并区被覆盖格要不要能带上自己的值，openpyxl 拒绝写入（MergedCell 只读）；\n"
            "- `w:vMerge` 元素在 XML 里有固定次序，python-docx 直接追加容易写出 Word 不认的顺序。\n\n"
            "运行方式见该目录 README。\n"
        )


def gen_09_docx():
    """OFF-09：首行有文本的表（headerRow=0）/ 首行全空的表（headerRow=null）"""
    directory = out_dir("09")

    filled = Document()
    table = filled.add_table(rows=2, cols=2)
    for col, value in enumerate(["评分项", "分值"]):
        table.cell(0, col).text = value
    for col, value in enumerate(["A1 报价", "30"]):
        table.cell(1, col).text = value
    filled.save(os.path.join(directory, "01-首行有文本的表.docx"))

    blank = Document()
    table = blank.add_table(rows=2, cols=2)
    table.cell(1, 0).text = "A1 报价"
    table.cell(1, 1).text = "30"
    blank.save(os.path.join(directory, "02-首行全空的表.docx"))


def gen_10_samples():
    """OFF-10：DOCX 页眉页脚 / XLSX 页眉三区与格式码"""
    directory = out_dir("10")

    doc = Document()
    doc.add_paragraph("正文一段")
    doc.sections[0].header.paragraphs[0].text = "文档页眉"
    doc.sections[0].footer.paragraphs[0].text = "文档页脚"
    doc.save(os.path.join(directory, "01-页眉页脚.docx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "页眉表"
    ws["A1"] = "A"
    ws["B1"] = "B"
    ws.oddHeader.left.text = "报表标题"
    ws.oddHeader.center.text = '&"Arial,Bold"机密'
    ws.oddHeader.right.text = "第 &P 页 && 共 &N 页"
    ws.oddFooter.center.text = "页脚文字"
    wb.save(os.path.join(directory, "02-页眉三区与格式码.xlsx"))


def gen_11_xlsx():
    """OFF-11：内嵌一张图片（IMAGE 元素 + 每表一条 OCR_IMAGE 事实）"""
    directory = out_dir("11")
    from openpyxl.drawing.image import Image as XlsxImage

    wb = Workbook()
    ws = wb.active
    ws.title = "图表"
    ws["A1"] = "指标"
    ws["B1"] = "值"
    image = XlsxImage(_png_file(directory))
    image.width, image.height = 20, 20
    ws.add_image(image, "D2")
    wb.save(os.path.join(directory, "01-内嵌一张图.xlsx"))


def gen_14_docx():
    """OFF-14：文本框文字（python-docx 不直接支持文本框，这里注入 w:txbxContent XML）"""
    directory = out_dir("14")
    doc = Document()
    doc.add_paragraph("正文一段")
    drawing = (
        '<w:p %s><w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="1828800" cy="457200"/><wp:docPr id="1" name="文本框 1"/>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1828800" cy="457200"/>'
        '</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
        '<wps:txbx><w:txbxContent><w:p><w:r><w:t>文本框里的文字</w:t></w:r></w:p></w:txbxContent></wps:txbx>'
        '<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
        % nsdecls("w", "wp", "a")
    )
    doc.element.body.append(parse_xml(drawing))
    doc.save(os.path.join(directory, "01-文本框文字.docx"))


def gen_15_note():
    """OFF-15：.doc（HWPF 旧格式）无法用 POI 写出，目录内放说明"""
    directory = out_dir("15")
    with open(os.path.join(directory, "说明-旧格式样例无法程序化生成.md"), "w", encoding="utf-8") as handle:
        handle.write(
            "# OFF-15 样例说明\n\n"
            "本条目需要一份带域代码（PAGE / HYPERLINK 等，控制符 0x13 / 0x14 / 0x15）的 **.doc（Word 97-2003 二进制）**。\n\n"
            "为什么这里没有文件：\n\n"
            "- POI 的 HWPF 只能读旧二进制格式、无法从零写出；\n"
            "- 当前环境 PATH 与 DSH 资源里都没有 LibreOffice，无法把 .docx 转成 .doc。\n\n"
            "手工准备方式（任选其一）：\n\n"
            "1. 用 Word 或 WPS 新建文档，插入「页码」域与「超链接」域，另存为 **Word 97-2003 文档（*.doc）**，放到本目录；\n"
            "2. 用 LibreOffice：`soffice --headless --convert-to doc 任意.docx`。\n\n"
            "期望结果（解析产物）：\n\n"
            "- 段落文本里**不应**出现域指令（`PAGE`、`HYPERLINK \"http://…\"`）与 0x13/0x14/0x15 控制符；\n"
            "- 域结果文本保留，例如「第 3 页」中的 `3`；\n"
            "- 单元格文本同样按该口径剥离。\n"
        )


def main():
    os.makedirs(SAMPLES, exist_ok=True)
    gen_01_pptx()
    gen_02_docx()
    gen_03_docx()
    gen_04_docx()
    gen_05_docx()
    gen_06_xlsx()
    gen_07_xlsx()
    gen_08_samples()
    gen_09_docx()
    gen_10_samples()
    gen_11_xlsx()
    gen_14_docx()
    gen_15_note()
    print("样例已生成到 %s" % os.path.join(SAMPLES, "office-parse-*"))


if __name__ == "__main__":
    main()
